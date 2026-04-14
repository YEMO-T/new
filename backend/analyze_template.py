#!/usr/bin/env python3
"""
分析模板结构 - 诊断非标准占位符问题
"""

import os
import sys
from pathlib import Path

backend_dir = Path(__file__).parent
sys.path.insert(0, str(backend_dir))

from pptx import Presentation

template_path = os.path.join(backend_dir, 'data', 'templates', '290eaaee-86be-4dea-a558-e10d2934df48.pptx')
prs = Presentation(template_path)

print("="*70)
print("模板结构深度分析")
print("="*70)
print(f"模板文件: {template_path}")
print(f"版式数量: {len(prs.slide_layouts)}")
print()

for idx, layout in enumerate(prs.slide_layouts):
    print(f"\n{'='*70}")
    print(f"版式 [{idx}]: {layout.name or '(unnamed)'}")
    print(f"{'='*70}")
    
    print(f"\n  占位符 (placeholders):")
    if hasattr(layout, 'placeholders') and list(layout.placeholders):
        for ph in layout.placeholders:
            try:
                ph_type = str(ph.placeholder_format.type)
                ph_idx = ph.placeholder_format.idx
                ph_name = ph.name or ''
                print(f"    - [{ph_idx}] {ph_name} (type={ph_type})")
            except Exception as e:
                print(f"    - (解析失败: {e})")
    else:
        print("    (无标准占位符)")
    
    print(f"\n  形状 (shapes):")
    shape_count = 0
    text_frame_count = 0
    for shape in layout.shapes:
        shape_count += 1
        has_tf = hasattr(shape, 'text_frame')
        if has_tf:
            text_frame_count += 1
        
        if shape_count <= 10:  # 只显示前10个
            shape_name = shape.name or '(unnamed)'
            shape_type = type(shape).__name__
            tf_info = "有text_frame" if has_tf else "无text_frame"
            
            try:
                text_preview = ""
                if has_tf and shape.text_frame.text:
                    text_preview = f" | 文本: '{shape.text_frame.text[:30]}...'"
            except:
                text_preview = ""
            
            print(f"    - {shape_name} ({shape_type}) [{tf_info}]{text_preview}")
    
    if shape_count > 10:
        print(f"    ... 还有 {shape_count - 10} 个形状")
    
    print(f"\n  统计: {shape_count}个形状, {text_frame_count}个带文本框")

print("\n" + "="*70)
print("结论与建议")
print("="*70)
print("""
这是一个自定义设计模板，特点：
1. 所有版式都没有标准占位符(TITLE/BODY等)
2. 内容通过普通形状(shape)和文本框(text_frame)实现
3. 版式名称如'标题幻灯片'、'自定义版式'等只是命名，不代表实际结构

对于这种模板的渲染策略：
- 方案A: 使用Blank版式 + add_textbox完全自定义布局
- 方案B: 选择一个版式后，遍历其shapes找到可写入的text_frame
- 方案C: 使用紧急注入模式(当前方案)，但优化减少日志噪音

推荐: 采用方案B，在版式中寻找可用的text_frame进行写入
""")