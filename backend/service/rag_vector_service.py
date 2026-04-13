"""
RAG向量化服务 - 基于 LangChain 的文本分块与向量化
支持 TXT/PDF/PPTX 文档解析
"""

import os
import logging
import asyncio
import time
from typing import List, Dict, Any, Optional, Tuple
from io import BytesIO

if not os.environ.get('HF_ENDPOINT'):
    os.environ['HF_ENDPOINT'] = 'https://hf-mirror.com'

from langchain_text_splitters import RecursiveCharacterTextSplitter

from repository.supabase_client import get_supabase_client

logger = logging.getLogger(__name__)

_embedding_model = None
_text_splitter = None

MODEL_CACHE_DIR = os.path.join(os.path.dirname(__file__), '..', 'models')
os.makedirs(MODEL_CACHE_DIR, exist_ok=True)


def get_embedding_model():
    """
    获取或初始化向量嵌入模型（单例模式）
    优先使用本地模型，否则尝试在线下载
    """
    global _embedding_model
    if _embedding_model is None:
        local_model_path = os.path.join(MODEL_CACHE_DIR, 'all-MiniLM-L6-v2')
        
        if os.path.exists(local_model_path):
            logger.info(f"[INFO] 使用本地模型: {local_model_path}")
            try:
                from langchain_community.embeddings import HuggingFaceEmbeddings
                _embedding_model = HuggingFaceEmbeddings(
                    model_name=local_model_path,
                    model_kwargs={'device': 'cpu'},
                    encode_kwargs={'normalize_embeddings': True}
                )
                logger.info("[OK] 本地模型加载完成")
                return _embedding_model
            except Exception as e:
                logger.warning(f"[WARN] 本地模型加载失败: {e}")
        
        logger.info("[INFO] 正在加载向量模型...")
        logger.info(f"[INFO] 使用镜像: {os.environ.get('HF_ENDPOINT', 'default')}")
        
        try:
            from langchain_community.embeddings import HuggingFaceEmbeddings
            _embedding_model = HuggingFaceEmbeddings(
                model_name="sentence-transformers/all-MiniLM-L6-v2",
                model_kwargs={'device': 'cpu', 'cache_folder': MODEL_CACHE_DIR},
                encode_kwargs={'normalize_embeddings': True}
            )
            logger.info("[OK] 向量模型加载完成")
        except Exception as e:
            logger.error(f"[ERR] 模型加载失败: {e}")
            logger.info("[INFO] 尝试使用备用方案...")
            _embedding_model = FallbackEmbedding()
    
    return _embedding_model


class FallbackEmbedding:
    """
    备用嵌入方案：使用简单的 TF-IDF 或哈希嵌入
    当 HuggingFace 模型无法加载时使用
    """
    
    def __init__(self):
        self.dimension = 384
        logger.warning("[WARN] 使用备用嵌入方案（简化版），检索效果可能降低")
    
    def embed_documents(self, texts: List[str]) -> List[List[float]]:
        return [self.embed_query(text) for text in texts]
    
    def embed_query(self, text: str) -> List[float]:
        import hashlib
        import numpy as np
        
        text_hash = hashlib.md5(text.encode()).hexdigest()
        np.random.seed(int(text_hash[:8], 16))
        embedding = np.random.randn(self.dimension).astype(np.float32)
        embedding = embedding / np.linalg.norm(embedding)
        return embedding.tolist()


def get_text_splitter():
    """
    获取或初始化文本分割器（单例模式）
    使用 LangChain 的 RecursiveCharacterTextSplitter
    """
    global _text_splitter
    if _text_splitter is None:
        _text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=512,
            chunk_overlap=100,
            length_function=len,
            separators=["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""]
        )
        logger.info("[OK] 文本分割器初始化完成")
    return _text_splitter


def parse_pptx_content(file_bytes: bytes) -> str:
    """
    解析 PPTX 文件内容为纯文本
    
    Args:
        file_bytes: PPTX 文件的二进制内容
    
    Returns:
        提取的文本内容
    """
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        
        prs = Presentation(BytesIO(file_bytes))
        text_parts = []
        
        for slide_idx, slide in enumerate(prs.slides, 1):
            slide_text = [f"\n=== 幻灯片 {slide_idx} ==="]
            
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        para_text = "".join([run.text for run in paragraph.runs])
                        if para_text.strip():
                            slide_text.append(para_text.strip())
                
                elif shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_text = []
                        for cell in row.cells:
                            if cell.text_frame:
                                cell_text = " ".join([
                                    run.text 
                                    for para in cell.text_frame.paragraphs 
                                    for run in para.runs
                                ])
                                if cell_text.strip():
                                    row_text.append(cell_text.strip())
                        if row_text:
                            slide_text.append(" | ".join(row_text))
            
            text_parts.append("\n".join(slide_text))
        
        full_text = "\n\n".join(text_parts)
        logger.info(f"[OK] PPTX解析完成: {len(prs.slides)} 页, {len(full_text)} 字符")
        return full_text
        
    except Exception as e:
        logger.error(f"[ERR] PPTX解析失败: {e}")
        return f"[PPTX解析错误] {str(e)}"


def parse_pdf_content(file_bytes: bytes) -> str:
    """
    解析 PDF 文件内容为纯文本
    
    Args:
        file_bytes: PDF 文件的二进制内容
    
    Returns:
        提取的文本内容
    """
    try:
        import PyPDF2
        
        reader = PyPDF2.PdfReader(BytesIO(file_bytes))
        text_parts = []
        
        for page_idx, page in enumerate(reader.pages, 1):
            page_text = page.extract_text()
            if page_text.strip():
                text_parts.append(f"[第{page_idx}页]\n{page_text}")
        
        full_text = "\n\n".join(text_parts)
        logger.info(f"[OK] PDF解析完成: {len(reader.pages)} 页, {len(full_text)} 字符")
        return full_text
        
    except Exception as e:
        logger.error(f"[ERR] PDF解析失败: {e}")
        return f"[PDF解析错误] {str(e)}"


def parse_document_content(file_bytes: bytes, file_type: str, original_name: str = "") -> Tuple[str, int]:
    """
    解析文档内容（支持 TXT/PDF/PPTX）
    
    Args:
        file_bytes: 文件二进制内容
        file_type: 文件类型（txt/pdf/pptx）
        original_name: 原始文件名
    
    Returns:
        (文本内容, 页数/段落数)
    """
    file_type = file_type.lower().strip('.')
    
    if file_type == 'pptx':
        text = parse_pptx_content(file_bytes)
        page_count = text.count("=== 幻灯片")
        return text, max(page_count, 1)
    
    elif file_type == 'pdf':
        text = parse_pdf_content(file_bytes)
        page_count = text.count("[第")
        return text, max(page_count, 1)
    
    else:
        try:
            text = file_bytes.decode('utf-8')
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode('gbk', errors='ignore')
            except:
                text = file_bytes.decode('utf-8', errors='ignore')
        
        return text, 1


def chunk_text(text: str) -> List[str]:
    """
    使用 LangChain 分割文本
    
    Args:
        text: 原始文本
    
    Returns:
        文本块列表
    """
    if not text or len(text.strip()) == 0:
        return []
    
    splitter = get_text_splitter()
    chunks = splitter.split_text(text)
    
    chunks = [chunk.strip() for chunk in chunks if chunk.strip()]
    
    logger.info(f"[INFO] 文本分块完成: {len(text)} 字符 -> {len(chunks)} 块")
    return chunks


def embed_text(text: str) -> Optional[List[float]]:
    """
    将文本转换为向量嵌入
    
    Args:
        text: 要向量化的文本
    
    Returns:
        384维向量列表，失败返回 None
    """
    try:
        model = get_embedding_model()
        embedding = model.embed_query(text)
        return embedding
    except Exception as e:
        logger.error(f"[ERR] 向量化失败: {e}")
        return None


async def vectorize_knowledge_item(
    item_id: str,
    content: str,
    source_resource: str = "unknown",
    page_number: int = 0
) -> Dict[str, Any]:
    """
    对单个知识库条目进行向量化处理
    
    Args:
        item_id: 知识库条目ID
        content: 文本内容
        source_resource: 来源资源名称
        page_number: 页码
    
    Returns:
        处理结果
    """
    start_time = time.time()
    
    try:
        supabase = get_supabase_client()
        
        supabase.table("knowledge_items").update({
            "vector_status": "processing"
        }).eq("id", item_id).execute()
        
        logger.info(f"[INFO] 开始向量化条目: {item_id}")
        
        chunks = chunk_text(content)
        if not chunks:
            raise ValueError("文本分块失败或内容为空")
        
        vectors_to_insert = []
        for chunk_index, chunk_text_content in enumerate(chunks):
            embedding = embed_text(chunk_text_content)
            if embedding is None:
                logger.warning(f"[WARN] 第{chunk_index}块向量化失败，跳过")
                continue
            
            vectors_to_insert.append({
                "knowledge_item_id": item_id,
                "chunk_index": chunk_index,
                "chunk_text": chunk_text_content,
                "vector_embedding": embedding,
                "source_resource": source_resource,
                "page_number": page_number
            })
        
        if vectors_to_insert:
            supabase.table("knowledge_vectors").insert(vectors_to_insert).execute()
            logger.info(f"[OK] 插入 {len(vectors_to_insert)} 个向量")
        
        processing_time = int((time.time() - start_time) * 1000)
        
        supabase.table("knowledge_items").update({
            "vector_status": "completed",
            "chunk_count": len(vectors_to_insert)
        }).eq("id", item_id).execute()
        
        try:
            supabase.table("vectorization_logs").insert({
                "knowledge_item_id": item_id,
                "status": "success",
                "chunk_count": len(vectors_to_insert),
                "processing_time_ms": processing_time
            }).execute()
        except Exception as log_err:
            logger.warning(f"[WARN] 日志记录失败: {log_err}")
        
        logger.info(f"[OK] 条目 {item_id} 向量化完成: {len(vectors_to_insert)} 块, 耗时 {processing_time}ms")
        
        return {
            "success": True,
            "chunk_count": len(vectors_to_insert),
            "item_id": item_id,
            "processing_time_ms": processing_time
        }
        
    except Exception as e:
        logger.error(f"[ERR] 向量化失败: {e}")
        
        try:
            supabase = get_supabase_client()
            supabase.table("knowledge_items").update({
                "vector_status": "failed"
            }).eq("id", item_id).execute()
        except Exception as update_err:
            logger.warning(f"[WARN] 更新状态失败: {update_err}")
        
        try:
            supabase = get_supabase_client()
            supabase.table("vectorization_logs").insert({
                "knowledge_item_id": item_id,
                "status": "failed",
                "error_message": str(e)[:500]
            }).execute()
        except Exception as log_error:
            logger.warning(f"[WARN] 记录失败日志异常: {log_error}")
        
        return {
            "success": False,
            "error": str(e),
            "item_id": item_id
        }


async def batch_vectorize_pending_items(user_id: Optional[str] = None) -> Dict[str, Any]:
    """
    批量处理所有待向量化的条目
    
    处理条件：
    - vector_status 为 'pending'
    - vector_status 为 NULL（未处理）
    - vector_status 为 'failed'（失败重试）
    
    Args:
        user_id: 用户ID（可选，不传则处理所有）
    
    Returns:
        处理结果统计
    """
    try:
        supabase = get_supabase_client()
        
        query = supabase.table("knowledge_items").select(
            "id, name, content"
        ).or_("vector_status.is.null,vector_status.eq.pending,vector_status.eq.failed")
        
        if user_id:
            query = query.eq("user_id", user_id)
        
        pending_items = query.execute()
        
        items = pending_items.data
        logger.info(f"[INFO] 发现 {len(items)} 个待向量化条目 (pending/null/failed)")
        
        results = {
            "total": len(items),
            "success": 0,
            "failed": 0,
            "items": []
        }
        
        for item in items:
            result = await vectorize_knowledge_item(
                item_id=item["id"],
                content=item.get("content", ""),
                source_resource=item.get("name", "unknown")
            )
            
            if result["success"]:
                results["success"] += 1
            else:
                results["failed"] += 1
            
            results["items"].append({
                "id": item["id"],
                "name": item.get("name"),
                "success": result["success"]
            })
            
            await asyncio.sleep(0.1)
        
        logger.info(f"[OK] 批量处理完成: 成功{results['success']}, 失败{results['failed']}")
        return results
        
    except Exception as e:
        logger.error(f"[ERR] 批量处理异常: {e}")
        return {"total": 0, "success": 0, "failed": 0, "error": str(e)}


async def search_rag(
    query: str,
    user_id: Optional[str] = None,
    grade: Optional[str] = None,
    subject: Optional[str] = None,
    top_k: int = 5,
    min_confidence: float = 0.5
) -> List[Dict[str, Any]]:
    """
    RAG 向量相似度检索
    
    Args:
        query: 查询文本
        user_id: 用户ID（用于权限过滤）
        grade: 年级过滤
        subject: 学科过滤
        top_k: 返回结果数
        min_confidence: 最小置信度
    
    Returns:
        检索结果列表
    """
    try:
        query_embedding = embed_text(query)
        if query_embedding is None:
            logger.error("[ERR] 查询向量化失败")
            return []
        
        logger.info(f"[INFO] 执行RAG搜索: query='{query}', user={user_id}, grade={grade}, subject={subject}")
        
        supabase = get_supabase_client()
        
        try:
            result = supabase.rpc(
                "rag_search",
                {
                    "query_embedding": query_embedding,
                    "query_user_id": user_id,
                    "query_grade": grade,
                    "query_subject": subject,
                    "match_threshold": min_confidence,
                    "match_count": top_k
                }
            ).execute()
            
            results = result.data if result.data else []
            
            logger.info(f"[OK] 检索完成: 返回 {len(results)} 条结果")
            
            try:
                supabase.table("search_history").insert({
                    "user_id": user_id,
                    "query": query,
                    "grade": grade,
                    "subject": subject,
                    "results_count": len(results),
                    "top_confidence_score": results[0]["similarity"] if results else 0.0
                }).execute()
            except Exception as e:
                logger.warning(f"[WARN] 记录搜索历史异常: {e}")
            
            return results
            
        except Exception as rpc_error:
            logger.warning(f"[WARN] RPC调用失败，使用Python实现: {rpc_error}")
            
            all_vectors = supabase.table("knowledge_vectors").select(
                "id, chunk_text, source_resource, page_number, chunk_index, vector_embedding, knowledge_item_id"
            ).execute()
            
            results_with_scores = []
            
            for vector_record in all_vectors.data:
                try:
                    stored_embedding = vector_record.get("vector_embedding")
                    if not stored_embedding:
                        continue
                    
                    similarity = _cosine_similarity(query_embedding, stored_embedding)
                    
                    if similarity >= min_confidence:
                        item_info = supabase.table("knowledge_items").select(
                            "id, name, visibility, grade_level, subject, user_id"
                        ).eq("id", vector_record["knowledge_item_id"]).single().execute()
                        
                        if item_info.data:
                            item = item_info.data
                            
                            if item["visibility"] != "public" and item["user_id"] != user_id:
                                continue
                            
                            if grade and item.get("grade_level") != grade:
                                continue
                            if subject and item.get("subject") != subject:
                                continue
                            
                            results_with_scores.append({
                                "id": vector_record["id"],
                                "knowledge_item_id": vector_record["knowledge_item_id"],
                                "chunk_text": vector_record["chunk_text"],
                                "chunk_index": vector_record["chunk_index"],
                                "source_resource": vector_record["source_resource"],
                                "page_number": vector_record["page_number"],
                                "similarity": float(similarity),
                                "item_name": item.get("name"),
                                "item_visibility": item.get("visibility"),
                                "item_grade": item.get("grade_level"),
                                "item_subject": item.get("subject")
                            })
                            
                except Exception as e:
                    logger.warning(f"[WARN] 处理向量记录异常: {e}")
                    continue
            
            results_with_scores.sort(key=lambda x: x["similarity"], reverse=True)
            final_results = results_with_scores[:top_k]
            
            logger.info(f"[OK] 检索完成(Python): 返回 {len(final_results)} 条结果")
            return final_results
        
    except Exception as e:
        logger.error(f"[ERR] RAG搜索异常: {e}")
        return []


def _parse_vector(vec) -> Optional[List[float]]:
    """
    健壮的向量解析函数
    支持多种格式：
    - List[float]: 直接返回
    - "0.1,0.2,0.3": 逗号分隔字符串
    - "0.1 0.2 0.3": 空格分隔字符串
    - "[0.1, 0.2, 0.3]": 带方括号的字符串
    - numpy数组: 转换为列表
    """
    if vec is None:
        return None
    
    if isinstance(vec, (list, tuple)):
        try:
            return [float(x) if x is not None else 0.0 for x in vec]
        except (ValueError, TypeError):
            return None
    
    if hasattr(vec, 'tolist'):
        try:
            return vec.tolist()
        except Exception:
            pass
    
    if not isinstance(vec, str):
        return None
    
    vec_str = vec.strip()
    if not vec_str:
        return None
    
    vec_str = vec_str.strip('[]').strip('()').strip('"').strip("'")
    vec_str = vec_str.replace('\n', ' ').replace('\t', ' ')
    
    separators = [',', ' ', ';', '|']
    parts = None
    
    for sep in separators:
        if sep in vec_str:
            parts = [p.strip() for p in vec_str.split(sep) if p.strip()]
            break
    
    if parts is None:
        parts = [vec_str] if vec_str else []
    
    result = []
    for part in parts:
        try:
            cleaned = part.strip().replace('"', '').replace("'", '')
            if cleaned:
                result.append(float(cleaned))
        except (ValueError, TypeError):
            continue
    
    return result if result else None


def _cosine_similarity(vec1: List[float], vec2) -> float:
    """
    计算两个向量的余弦相似度
    支持多种向量格式输入
    """
    try:
        import numpy as np
        
        parsed_vec1 = _parse_vector(vec1)
        parsed_vec2 = _parse_vector(vec2)
        
        if parsed_vec1 is None or parsed_vec2 is None:
            logger.error(f"[ERR] 向量解析失败: vec1={type(vec1)}, vec2={type(vec2)}")
            return 0.0
        
        if len(parsed_vec1) == 0 or len(parsed_vec2) == 0:
            logger.error("[ERR] 向量为空")
            return 0.0
        
        arr1 = np.array(parsed_vec1, dtype=np.float64)
        arr2 = np.array(parsed_vec2, dtype=np.float64)
        
        if arr1.shape != arr2.shape:
            min_len = min(len(arr1), len(arr2))
            if min_len == 0:
                return 0.0
            arr1 = arr1[:min_len]
            arr2 = arr2[:min_len]
            logger.debug(f"[DEBUG] 向量维度对齐: {min_len}")
        
        norm1 = np.linalg.norm(arr1)
        norm2 = np.linalg.norm(arr2)
        
        if norm1 == 0 or norm2 == 0:
            return 0.0
        
        dot_product = np.dot(arr1, arr2)
        similarity = dot_product / (norm1 * norm2)
        
        return float(similarity)
        
    except Exception as e:
        logger.error(f"[ERR] 相似度计算异常: {e}")
        return 0.0


async def get_vectorization_status(item_id: str) -> Dict[str, Any]:
    """
    获取条目的向量化状态
    """
    try:
        supabase = get_supabase_client()
        
        item = supabase.table("knowledge_items").select(
            "id, name, vector_status, chunk_count"
        ).eq("id", item_id).execute()
        
        if not item.data:
            return {"error": "条目不存在"}
        
        return {
            "item_id": item_id,
            "name": item.data[0]["name"],
            "status": item.data[0]["vector_status"],
            "chunk_count": item.data[0]["chunk_count"]
        }
        
    except Exception as e:
        logger.error(f"[ERR] 获取状态异常: {e}")
        return {"error": str(e)}


def delete_knowledge_vectors(item_id: str) -> bool:
    """
    删除指定知识库条目的所有向量
    """
    try:
        supabase = get_supabase_client()
        supabase.table("knowledge_vectors").delete().eq("knowledge_item_id", item_id).execute()
        logger.info(f"[OK] 已删除条目 {item_id} 的所有向量")
        return True
    except Exception as e:
        logger.error(f"[ERR] 删除向量失败: {e}")
        return False
