"""
PPT生成器 - 基于模板生成样式完全一致的PPT
核心功能：克隆母版、复用版式、样式继承、内容填充
"""

import io
import os
import logging
import tempfile
import copy
from typing import Dict, List, Any, Optional, Tuple
from datetime import datetime

from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn, nsmap
from pptx.oxml import parse_xml
from lxml import etree

logger = logging.getLogger(__name__)


class PPTGenerator:
    """
    PPT生成器 - 基于模板生成同风格PPT
    
    核心原理：
    1. 加载模板PPT，保留所有母版和版式
    2. 清空模板中的示例幻灯片
    3. 基于模板版式创建新幻灯片
    4. 向占位符填充内容，自动继承模板样式
    """
    
    def __init__(self, template_path: str):
        """
        初始化生成器，加载模板
        
        Args:
            template_path: 模板文件路径
        """
        self.template_path = template_path
        self.prs: Optional[Presentation] = None
        self.slide_layouts: Dict[str, Any] = {}
        self.master_styles: Dict[str, Any] = {}
        
        self._load_template()
    
    def _load_template(self):
        """加载模板并提取样式信息"""
        try:
            if not os.path.exists(self.template_path):
                raise FileNotFoundError(f"模板文件不存在: {self.template_path}")
            
            self.prs = Presentation(self.template_path)
            logger.info(f"[PPTGenerator] 加载模板成功: {self.template_path}")
            logger.info(f"[PPTGenerator] 母版数量: {len(self.prs.slide_masters)}")
            logger.info(f"[PPTGenerator] 版式数量: {len(self.prs.slide_layouts)}")
            
            self._extract_layouts_info()
            self._extract_master_styles()
            
            self._clear_template_slides()
            
        except Exception as e:
            logger.error(f"[PPTGenerator] 加载模板失败: {type(e).__name__}: {e}")
            raise
    
    def _extract_layouts_info(self):
        """提取所有版式信息"""
        for idx, layout in enumerate(self.prs.slide_layouts):
            layout_info = {
                'index': idx,
                'name': layout.name,
                'placeholders': [],
                'has_title': False,
                'has_body': False,
            }
            
            for shape in layout.placeholders:
                ph_type = self._get_placeholder_type(shape)
                ph_info = {
                    'idx': shape.placeholder_format.idx,
                    'type': ph_type,
                    'name': shape.name,
                    'left': shape.left,
                    'top': shape.top,
                    'width': shape.width,
                    'height': shape.height,
                }
                layout_info['placeholders'].append(ph_info)
                
                if ph_type == 'title':
                    layout_info['has_title'] = True
                elif ph_type == 'body':
                    layout_info['has_body'] = True
            
            self.slide_layouts[layout.name] = layout_info
            self.slide_layouts[str(idx)] = layout_info
            
            logger.debug(f"[PPTGenerator] 版式 [{idx}] {layout.name}: "
                        f"title={layout_info['has_title']}, body={layout_info['has_body']}, "
                        f"placeholders={len(layout_info['placeholders'])}")
    
    def _get_placeholder_type(self, shape) -> str:
        """获取占位符类型"""
        try:
            ph_format = shape.placeholder_format
            ph_type = ph_format.type
            
            type_map = {
                PP_PLACEHOLDER.TITLE: 'title',
                PP_PLACEHOLDER.CENTER_TITLE: 'title',
                PP_PLACEHOLDER.SUBTITLE: 'subtitle',
                PP_PLACEHOLDER.BODY: 'body',
                PP_PLACEHOLDER.OBJECT: 'body',
                PP_PLACEHOLDER.VERTICAL_BODY: 'body',
                PP_PLACEHOLDER.VERTICAL_TITLE: 'title',
            }
            
            return type_map.get(ph_type, 'other')
        except:
            return 'other'
    
    def _extract_master_styles(self):
        """提取母版样式信息"""
        try:
            for master_idx, master in enumerate(self.prs.slide_masters):
                master_info = {
                    'index': master_idx,
                    'name': getattr(master, 'name', f'Master_{master_idx}'),
                    'title_style': self._extract_text_style_from_master(master, 'title'),
                    'body_style': self._extract_text_style_from_master(master, 'body'),
                }
                self.master_styles[f'master_{master_idx}'] = master_info
                
            logger.info(f"[PPTGenerator] 提取了 {len(self.master_styles)} 个母版样式")
        except Exception as e:
            logger.warning(f"[PPTGenerator] 提取母版样式失败: {e}")
    
    def _extract_text_style_from_master(self, master, style_type: str) -> Dict[str, Any]:
        """从母版提取文本样式"""
        style_info = {
            'font_name': None,
            'font_size': None,
            'font_bold': None,
            'font_color': None,
            'alignment': None,
        }
        
        try:
            for shape in master.shapes:
                if not hasattr(shape, 'text_frame'):
                    continue
                
                is_title = self._is_title_shape(shape)
                
                if (style_type == 'title' and is_title) or (style_type == 'body' and not is_title):
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.font.name:
                            style_info['font_name'] = paragraph.font.name
                        if paragraph.font.size:
                            style_info['font_size'] = paragraph.font.size.pt
                        if paragraph.font.bold is not None:
                            style_info['font_bold'] = paragraph.font.bold
                        try:
                            if paragraph.font.color.rgb:
                                style_info['font_color'] = str(paragraph.font.color.rgb)
                        except:
                            pass
                        if paragraph.alignment is not None:
                            style_info['alignment'] = paragraph.alignment
                        break
                    break
        except Exception as e:
            logger.debug(f"[PPTGenerator] 提取样式失败: {e}")
        
        return style_info
    
    def _is_title_shape(self, shape) -> bool:
        """判断是否为标题形状"""
        try:
            if hasattr(shape, 'placeholder_format'):
                ph_type = shape.placeholder_format.type
                if ph_type in [PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE]:
                    return True
            
            if hasattr(shape, 'name'):
                name_lower = shape.name.lower()
                if 'title' in name_lower or '标题' in shape.name:
                    return True
            
            return False
        except:
            return False
    
    def _clear_template_slides(self):
        """清空模板中的示例幻灯片，保留母版"""
        try:
            slide_count = len(self.prs.slides)
            while len(self.prs.slides) > 0:
                rId = self.prs.slides._sldIdLst[0].rId
                self.prs.part.drop_rel(rId)
                del self.prs.slides._sldIdLst[0]
            
            logger.info(f"[PPTGenerator] 已清空 {slide_count} 张模板幻灯片")
        except Exception as e:
            logger.warning(f"[PPTGenerator] 清空幻灯片失败: {e}")
    
    def find_best_layout(self, slide_type: str = 'content') -> int:
        """
        根据幻灯片类型找到最合适的版式索引
        
        Args:
            slide_type: 幻灯片类型 ('cover', 'content', 'summary', 'ending')
            
        Returns:
            版式索引
        """
        type_keywords = {
            'cover': ['title slide', '封面', '标题页', 'title', 'cover'],
            'content': ['title and content', '正文', '标题和内容', 'content', 'two content'],
            'summary': ['section header', '总结', '结束', 'summary', 'closing', 'ending'],
            'ending': ['blank', '空白', 'ending', 'closing', 'thank'],
        }
        
        keywords = type_keywords.get(slide_type, type_keywords['content'])
        
        for idx, layout in enumerate(self.prs.slide_layouts):
            layout_name_lower = layout.name.lower()
            for keyword in keywords:
                if keyword in layout_name_lower:
                    logger.info(f"[PPTGenerator] 为类型 '{slide_type}' 找到版式: [{idx}] {layout.name}")
                    return idx
        
        if slide_type == 'cover':
            for idx, layout in enumerate(self.prs.slide_layouts):
                layout_info = self.slide_layouts.get(str(idx), {})
                if layout_info.get('has_title') and not layout_info.get('has_body'):
                    logger.info(f"[PPTGenerator] 为类型 '{slide_type}' 找到标题版式: [{idx}] {layout.name}")
                    return idx
        else:
            for idx, layout in enumerate(self.prs.slide_layouts):
                layout_info = self.slide_layouts.get(str(idx), {})
                if layout_info.get('has_title') and layout_info.get('has_body'):
                    logger.info(f"[PPTGenerator] 为类型 '{slide_type}' 找到内容版式: [{idx}] {layout.name}")
                    return idx
        
        default_idx = 0 if slide_type == 'cover' else min(1, len(self.prs.slide_layouts) - 1)
        logger.info(f"[PPTGenerator] 使用默认版式: [{default_idx}]")
        return default_idx
    
    def add_slide(self, slide_data: Dict[str, Any]) -> bool:
        """
        添加一张幻灯片
        
        Args:
            slide_data: 幻灯片数据
                - type: 幻灯片类型 ('cover', 'content', 'summary', 'ending')
                - title: 标题文本
                - content: 正文内容（字符串或字符串数组）
                - subtitle: 副标题
                - bullets: 要点列表
                
        Returns:
            是否成功
        """
        try:
            slide_type = slide_data.get('type', 'content')
            layout_idx = self.find_best_layout(slide_type)
            slide_layout = self.prs.slide_layouts[layout_idx]
            
            slide = self.prs.slides.add_slide(slide_layout)
            
            title_text = slide_data.get('title', '')
            subtitle_text = slide_data.get('subtitle', '')
            content_text = slide_data.get('content', '')
            bullets = slide_data.get('bullets', [])
            
            if slide.shapes.title:
                slide.shapes.title.text = title_text
                logger.debug(f"[PPTGenerator] 填充标题: {title_text[:30]}...")
            
            if subtitle_text:
                for shape in slide.placeholders:
                    ph_type = self._get_placeholder_type(shape)
                    if ph_type == 'subtitle':
                        shape.text = subtitle_text
                        logger.debug(f"[PPTGenerator] 填充副标题: {subtitle_text[:30]}...")
                        break
            
            body_placeholder = None
            for shape in slide.placeholders:
                ph_type = self._get_placeholder_type(shape)
                if ph_type == 'body':
                    body_placeholder = shape
                    break
            
            if body_placeholder:
                if bullets and isinstance(bullets, list):
                    self._fill_bullets(body_placeholder, bullets)
                    logger.debug(f"[PPTGenerator] 填充 {len(bullets)} 个要点")
                elif content_text:
                    body_placeholder.text = content_text
                    logger.debug(f"[PPTGenerator] 填充正文: {content_text[:30]}...")
            
            logger.info(f"[PPTGenerator] 添加幻灯片成功: type={slide_type}, title={title_text[:20]}...")
            return True
            
        except Exception as e:
            logger.error(f"[PPTGenerator] 添加幻灯片失败: {type(e).__name__}: {e}")
            return False
    
    def _fill_bullets(self, placeholder, bullets: List[str]):
        """向占位符填充要点列表"""
        try:
            text_frame = placeholder.text_frame
            text_frame.clear()
            
            for i, bullet in enumerate(bullets):
                if i == 0:
                    p = text_frame.paragraphs[0]
                else:
                    p = text_frame.add_paragraph()
                
                p.text = bullet
                p.level = 0
                
        except Exception as e:
            logger.warning(f"[PPTGenerator] 填充要点失败: {e}")
            placeholder.text = "\n".join(bullets)
    
    def generate(self, slides_data: List[Dict[str, Any]]) -> bytes:
        """
        生成PPT文件
        
        Args:
            slides_data: 幻灯片数据列表
            
        Returns:
            PPT文件的字节流
        """
        try:
            logger.info(f"[PPTGenerator] 开始生成PPT，共 {len(slides_data)} 页")
            
            for idx, slide_data in enumerate(slides_data):
                logger.info(f"[PPTGenerator] 处理第 {idx + 1} 页: {slide_data.get('type', 'content')}")
                self.add_slide(slide_data)
            
            output = io.BytesIO()
            self.prs.save(output)
            output.seek(0)
            
            result_bytes = output.getvalue()
            logger.info(f"[PPTGenerator] PPT生成成功，大小: {len(result_bytes)} 字节")
            
            return result_bytes
            
        except Exception as e:
            logger.error(f"[PPTGenerator] 生成PPT失败: {type(e).__name__}: {e}", exc_info=True)
            raise
    
    def save_to_file(self, output_path: str) -> bool:
        """
        保存PPT到文件
        
        Args:
            output_path: 输出文件路径
            
        Returns:
            是否成功
        """
        try:
            self.prs.save(output_path)
            logger.info(f"[PPTGenerator] PPT已保存到: {output_path}")
            return True
        except Exception as e:
            logger.error(f"[PPTGenerator] 保存PPT失败: {e}")
            return False


def generate_ppt_from_template(
    template_path: str,
    slides_data: List[Dict[str, Any]]
) -> bytes:
    """
    基于模板生成PPT的便捷函数
    
    Args:
        template_path: 模板文件路径
        slides_data: 幻灯片数据列表
        
    Returns:
        PPT文件的字节流
    """
    generator = PPTGenerator(template_path)
    return generator.generate(slides_data)


def get_template_local_path(template_id: str) -> str:
    """
    获取模板的本地路径
    
    Args:
        template_id: 模板ID
        
    Returns:
        模板文件路径
    """
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base_dir, 'data', 'templates', f"{template_id}.pptx")


def generate_pptx(slides: List[Any], template_id: Optional[str] = None) -> io.BytesIO:
    """
    将 PPTSlide 列表转换为 PPTX 文件流
    
    Args:
        slides: PPTSlide 对象列表
        template_id: 可选的模板ID
        
    Returns:
        PPTX 文件的字节流
    """
    from utils.pptx_builder import create_pptx_from_ai_json
    
    slides_data = []
    for slide in slides:
        slide_dict = {
            "title": slide.title,
            "content": slide.content,
            "page_type": slide.page_type,
            "layout_suggestion": slide.layout_suggestion or "bullet_points"
        }
        slides_data.append(slide_dict)
    
    return create_pptx_from_ai_json(slides_data, template_id)
