"""
PPT 样式生成器
基于提取的样式元数据，生成与模板风格一致的新 PPT
"""
import io
import os
import logging
from typing import Dict, Any, List, Optional
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

from utils.ppt_style_extractor import extract_ppt_style_metadata

logger = logging.getLogger(__name__)


class PPTStyleGenerator:
    """
    基于模板样式的 PPT 生成器
    复用模板的字体、颜色、布局，生成风格一致的新 PPT
    """
    
    def __init__(self, template_bytes: bytes = None, template_path: str = None):
        """
        初始化生成器
        
        Args:
            template_bytes: 模板文件的字节流
            template_path: 模板文件路径
        """
        self.prs = None
        self.style_metadata = None
        self.extracted_styles = {}
        
        if template_bytes:
            self._load_from_bytes(template_bytes)
        elif template_path and os.path.exists(template_path):
            self._load_from_path(template_path)
        else:
            self.prs = Presentation()
            logger.info("使用空白模板初始化")
    
    def _load_from_bytes(self, template_bytes: bytes):
        """从字节流加载模板"""
        try:
            self.prs = Presentation(io.BytesIO(template_bytes))
            self.style_metadata = extract_ppt_style_metadata(template_bytes)
            self.extracted_styles = self.style_metadata.get("extracted_styles", {})
            self._clear_slides()
            logger.info(f"从字节流加载模板成功，提取样式: {self.style_metadata.get('success')}")
        except Exception as e:
            logger.error(f"加载模板字节流失败: {e}")
            self.prs = Presentation()
    
    def _load_from_path(self, template_path: str):
        """从文件路径加载模板"""
        try:
            with open(template_path, 'rb') as f:
                template_bytes = f.read()
            self._load_from_bytes(template_bytes)
            logger.info(f"从文件加载模板: {template_path}")
        except Exception as e:
            logger.error(f"加载模板文件失败: {e}")
            self.prs = Presentation()
    
    def _clear_slides(self):
        """清除模板中的幻灯片，保留母版"""
        try:
            while len(self.prs.slides) > 0:
                rId = self.prs.slides._sldIdLst[0].rId
                self.prs.part.drop_rel(rId)
                del self.prs.slides._sldIdLst[0]
        except Exception as e:
            logger.warning(f"清除幻灯片失败: {e}")
    
    def _get_font_style(self, style_type: str = "body") -> Dict[str, Any]:
        """获取字体样式"""
        fonts = self.extracted_styles.get("fonts", {})
        return fonts.get(style_type, fonts.get("body", {
            "name": "微软雅黑",
            "size": 18,
            "color": "#333333",
            "bold": False,
            "italic": False
        }))
    
    def _get_color_style(self) -> Dict[str, str]:
        """获取颜色样式"""
        return self.extracted_styles.get("colors", {
            "primary": "#4472C4",
            "secondary": "#ED7D31",
            "background": "#FFFFFF",
            "text": "#333333",
            "accent": "#4472C4"
        })
    
    def _apply_font_to_run(self, run, font_style: Dict[str, Any]):
        """应用字体样式到 run 对象"""
        try:
            font = run.font
            
            if font_style.get("name"):
                font.name = font_style["name"]
                try:
                    rPr = run.element.get_or_add_rPr()
                    rPr.set(qn('a:latin'), font_style["name"])
                    rPr.set(qn('a:ea'), font_style["name"])
                except:
                    pass
            
            if font_style.get("size"):
                font.size = Pt(font_style["size"])
            
            if font_style.get("bold") is not None:
                font.bold = font_style["bold"]
            
            if font_style.get("italic") is not None:
                font.italic = font_style["italic"]
            
            if font_style.get("color"):
                color_hex = font_style["color"].lstrip('#')
                font.color.rgb = RGBColor(
                    int(color_hex[0:2], 16),
                    int(color_hex[2:4], 16),
                    int(color_hex[4:6], 16)
                )
        except Exception as e:
            logger.debug(f"应用字体样式失败: {e}")
    
    def _apply_font_to_paragraph(self, paragraph, font_style: Dict[str, Any]):
        """应用字体样式到段落"""
        if not paragraph.text:
            return
        
        for run in paragraph.runs:
            self._apply_font_to_run(run, font_style)
    
    def _find_layout_by_type(self, slide_type: str) -> int:
        """根据幻灯片类型查找合适的布局索引"""
        layouts = self.prs.slide_layouts
        
        layout_keywords = {
            'cover': ['title slide', '封面', '标题幻灯片', 'title'],
            'content': ['title and content', '标题和内容', '正文', 'content', 'two content'],
            'toc': ['table of contents', '目录', 'toc'],
            'ending': ['section header', '结束', '总结', 'ending', 'closing', 'blank']
        }
        
        keywords = layout_keywords.get(slide_type.lower(), layout_keywords['content'])
        
        for i, layout in enumerate(layouts):
            layout_name = getattr(layout, 'name', '').lower()
            for keyword in keywords:
                if keyword in layout_name:
                    return i
        
        if slide_type == 'cover':
            return 0
        return min(1, len(layouts) - 1)
    
    def add_cover_slide(self, title: str, subtitle: str = "") -> bool:
        """
        添加封面页
        
        Args:
            title: 主标题
            subtitle: 副标题
            
        Returns:
            是否成功
        """
        try:
            layout_idx = self._find_layout_by_type('cover')
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[layout_idx])
            
            title_style = self._get_font_style("title")
            subtitle_style = self._get_font_style("subtitle")
            
            if slide.shapes.title:
                title_shape = slide.shapes.title
                title_shape.text = title
                if title_shape.text_frame.paragraphs:
                    self._apply_font_to_paragraph(title_shape.text_frame.paragraphs[0], title_style)
            
            for shape in slide.placeholders:
                if shape == slide.shapes.title:
                    continue
                if hasattr(shape, 'text_frame') and shape.placeholder_format.idx == 1:
                    shape.text = subtitle
                    if shape.text_frame.paragraphs:
                        self._apply_font_to_paragraph(shape.text_frame.paragraphs[0], subtitle_style)
                    break
            
            logger.info(f"添加封面页: {title}")
            return True
            
        except Exception as e:
            logger.error(f"添加封面页失败: {e}")
            return False
    
    def add_content_slide(self, title: str, content: List[str], slide_type: str = "content") -> bool:
        """
        添加内容页
        
        Args:
            title: 页面标题
            content: 内容列表（每个元素为一个要点）
            slide_type: 页面类型
            
        Returns:
            是否成功
        """
        try:
            layout_idx = self._find_layout_by_type(slide_type)
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[layout_idx])
            
            title_style = self._get_font_style("title")
            body_style = self._get_font_style("body")
            
            if slide.shapes.title:
                title_shape = slide.shapes.title
                title_shape.text = title
                if title_shape.text_frame.paragraphs:
                    self._apply_font_to_paragraph(title_shape.text_frame.paragraphs[0], title_style)
            
            body_placeholder = None
            for shape in slide.placeholders:
                if shape == slide.shapes.title:
                    continue
                ph_type = getattr(shape.placeholder_format, 'type', None)
                if ph_type in [PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT]:
                    body_placeholder = shape
                    break
            
            if not body_placeholder:
                for shape in slide.placeholders:
                    if shape != slide.shapes.title and hasattr(shape, 'text_frame'):
                        body_placeholder = shape
                        break
            
            if body_placeholder and content:
                tf = body_placeholder.text_frame
                tf.clear()
                
                for i, point in enumerate(content):
                    if i == 0:
                        p = tf.paragraphs[0]
                    else:
                        p = tf.add_paragraph()
                    
                    p.text = str(point)
                    p.level = 0
                    p.space_before = Pt(6)
                    p.space_after = Pt(6)
                    
                    self._apply_font_to_paragraph(p, body_style)
            
            logger.info(f"添加内容页: {title}, {len(content)} 个要点")
            return True
            
        except Exception as e:
            logger.error(f"添加内容页失败: {e}")
            return False
    
    def add_toc_slide(self, title: str, items: List[str]) -> bool:
        """
        添加目录页
        
        Args:
            title: 目录标题
            items: 目录项列表
            
        Returns:
            是否成功
        """
        try:
            layout_idx = self._find_layout_by_type('toc')
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[layout_idx])
            
            title_style = self._get_font_style("title")
            body_style = self._get_font_style("body")
            
            if slide.shapes.title:
                title_shape = slide.shapes.title
                title_shape.text = title
                if title_shape.text_frame.paragraphs:
                    self._apply_font_to_paragraph(title_shape.text_frame.paragraphs[0], title_style)
            
            body_placeholder = None
            for shape in slide.placeholders:
                if shape == slide.shapes.title:
                    continue
                if hasattr(shape, 'text_frame'):
                    body_placeholder = shape
                    break
            
            if body_placeholder and items:
                tf = body_placeholder.text_frame
                tf.clear()
                
                for i, item in enumerate(items):
                    if i == 0:
                        p = tf.paragraphs[0]
                    else:
                        p = tf.add_paragraph()
                    
                    p.text = f"{i + 1}. {item}"
                    p.level = 0
                    p.space_before = Pt(8)
                    p.space_after = Pt(8)
                    
                    self._apply_font_to_paragraph(p, body_style)
            
            logger.info(f"添加目录页: {title}, {len(items)} 项")
            return True
            
        except Exception as e:
            logger.error(f"添加目录页失败: {e}")
            return False
    
    def add_ending_slide(self, title: str = "感谢观看", subtitle: str = "Q & A") -> bool:
        """
        添加结束页
        
        Args:
            title: 结束语
            subtitle: 副标题
            
        Returns:
            是否成功
        """
        try:
            layout_idx = self._find_layout_by_type('ending')
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[layout_idx])
            
            title_style = self._get_font_style("title")
            subtitle_style = self._get_font_style("subtitle")
            
            if slide.shapes.title:
                title_shape = slide.shapes.title
                title_shape.text = title
                if title_shape.text_frame.paragraphs:
                    p = title_shape.text_frame.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    self._apply_font_to_paragraph(p, title_style)
            
            for shape in slide.placeholders:
                if shape == slide.shapes.title:
                    continue
                if hasattr(shape, 'text_frame'):
                    shape.text = subtitle
                    if shape.text_frame.paragraphs:
                        p = shape.text_frame.paragraphs[0]
                        p.alignment = PP_ALIGN.CENTER
                        self._apply_font_to_paragraph(p, subtitle_style)
                    break
            
            logger.info(f"添加结束页: {title}")
            return True
            
        except Exception as e:
            logger.error(f"添加结束页失败: {e}")
            return False
    
    def generate_from_slides_data(self, slides_data: List[Dict[str, Any]]) -> bool:
        """
        根据幻灯片数据列表生成完整 PPT
        
        Args:
            slides_data: 幻灯片数据列表，每个元素包含:
                - type: 页面类型 (cover/content/toc/ending)
                - title: 标题
                - content: 内容列表或字符串
                - subtitle: 副标题（可选）
                
        Returns:
            是否成功
        """
        if not slides_data:
            logger.warning("没有幻灯片数据")
            return False
        
        success_count = 0
        
        for i, slide_data in enumerate(slides_data):
            slide_type = slide_data.get("type", "content").lower()
            title = slide_data.get("title", "")
            content = slide_data.get("content", [])
            subtitle = slide_data.get("subtitle", "")
            
            if isinstance(content, str):
                content = [content] if content else []
            
            if i == 0 and slide_type != "cover":
                slide_type = "cover"
            
            if slide_type == "cover":
                if self.add_cover_slide(title, subtitle or (content[0] if content else "")):
                    success_count += 1
            elif slide_type == "toc":
                if self.add_toc_slide(title, content):
                    success_count += 1
            elif slide_type == "ending":
                if self.add_ending_slide(title, subtitle):
                    success_count += 1
            else:
                if self.add_content_slide(title, content, slide_type):
                    success_count += 1
        
        logger.info(f"生成 PPT 完成: {success_count}/{len(slides_data)} 页")
        return success_count > 0
    
    def build_stream(self) -> io.BytesIO:
        """
        生成 PPT 文件字节流
        
        Returns:
            PPT 文件的 BytesIO 对象
        """
        output = io.BytesIO()
        self.prs.save(output)
        output.seek(0)
        return output
    
    def get_style_metadata(self) -> Dict[str, Any]:
        """获取当前使用的样式元数据"""
        return self.style_metadata or {}


def generate_ppt_with_style(
    template_bytes: bytes,
    slides_data: List[Dict[str, Any]]
) -> io.BytesIO:
    """
    根据模板样式生成新 PPT
    
    Args:
        template_bytes: 模板文件的字节流
        slides_data: 幻灯片数据列表
        
    Returns:
        生成的 PPT 文件字节流
    """
    generator = PPTStyleGenerator(template_bytes=template_bytes)
    generator.generate_from_slides_data(slides_data)
    return generator.build_stream()
