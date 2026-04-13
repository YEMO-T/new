"""
PPTX构建器 - 增强版
用于将AI内容填入模板，保留母版设计和样式
"""

import io
import os
import logging
from typing import List, Dict, Optional, Union, Any
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

logger = logging.getLogger(__name__)


class PPTXBuilder:
    """
    增强版PPTX构建器
    
    核心功能：
    1. 加载模板，保留母版设计
    2. 清空模板示例幻灯片
    3. 提取并继承占位符样式
    4. 智能版式匹配
    """
    
    def __init__(self, template_path: Optional[str] = None):
        self.template_path = template_path
        self.slide_layouts_info: Dict[str, Any] = {}
        
        try:
            if template_path and os.path.exists(template_path):
                self.prs = Presentation(template_path)
                logger.info(f"[PPTXBuilder] 已载入模板: {os.path.basename(template_path)}")
                logger.info(f"[PPTXBuilder] 母版数量: {len(self.prs.slide_masters)}, 版式数量: {len(self.prs.slide_layouts)}")
                self._extract_layouts_info()
                self._clear_slides()
            else:
                self.prs = Presentation()
                logger.info("[PPTXBuilder] 已载入默认空白模板")
        except Exception as e:
            logger.error(f"[PPTXBuilder] 初始化失败: {e}")
            self.prs = Presentation()
    
    def _extract_layouts_info(self):
        """提取版式信息"""
        for idx, layout in enumerate(self.prs.slide_layouts):
            layout_info = {
                'index': idx,
                'name': layout.name,
                'has_title': False,
                'has_body': False,
                'placeholders': []
            }
            
            for shape in layout.placeholders:
                ph_type = self._get_placeholder_type(shape)
                layout_info['placeholders'].append({
                    'idx': shape.placeholder_format.idx,
                    'type': ph_type,
                    'name': shape.name
                })
                
                if ph_type == 'title':
                    layout_info['has_title'] = True
                elif ph_type == 'body':
                    layout_info['has_body'] = True
            
            self.slide_layouts_info[str(idx)] = layout_info
            self.slide_layouts_info[layout.name] = layout_info
    
    def _get_placeholder_type(self, shape) -> str:
        """获取占位符类型"""
        try:
            ph_type = shape.placeholder_format.type
            type_map = {
                PP_PLACEHOLDER.TITLE: 'title',
                PP_PLACEHOLDER.CENTER_TITLE: 'title',
                PP_PLACEHOLDER.SUBTITLE: 'subtitle',
                PP_PLACEHOLDER.BODY: 'body',
                PP_PLACEHOLDER.OBJECT: 'body',
                PP_PLACEHOLDER.VERTICAL_BODY: 'body',
                PP_PLACEHOLDER.VERTICAL_TITLE: 'title',
            }
            return type_map.get(ph_type, 'other')
        except:
            return 'other'
    
    def _clear_slides(self):
        """移除模板中的示例幻灯片，保留母版"""
        slide_count = len(self.prs.slides)
        while len(self.prs.slides) > 0:
            rId = self.prs.slides._sldIdLst[0].rId
            self.prs.part.drop_rel(rId)
            del self.prs.slides._sldIdLst[0]
        logger.info(f"[PPTXBuilder] 已清空 {slide_count} 张模板幻灯片")
    
    def _find_suitable_layout(self, slide_type: str = 'content') -> int:
        """根据幻灯片类型智能匹配版式"""
        type_keywords = {
            'cover': ['title slide', '封面', '标题页', 'title', 'cover'],
            'content': ['title and content', '正文', '标题和内容', 'content', 'two content'],
            'summary': ['section header', '总结', '结束', 'summary', 'closing'],
            'ending': ['blank', '空白', 'ending', 'closing', 'thank', '致谢'],
        }
        
        keywords = type_keywords.get(slide_type, type_keywords['content'])
        
        for idx, layout in enumerate(self.prs.slide_layouts):
            layout_name_lower = layout.name.lower()
            for keyword in keywords:
                if keyword in layout_name_lower:
                    logger.debug(f"[PPTXBuilder] 为 '{slide_type}' 找到版式: [{idx}] {layout.name}")
                    return idx
        
        if slide_type == 'cover':
            for idx, layout in enumerate(self.prs.slide_layouts):
                info = self.slide_layouts_info.get(str(idx), {})
                if info.get('has_title') and not info.get('has_body'):
                    return idx
        else:
            for idx, layout in enumerate(self.prs.slide_layouts):
                info = self.slide_layouts_info.get(str(idx), {})
                if info.get('has_title') and info.get('has_body'):
                    return idx
        
        return 0 if slide_type == 'cover' else min(1, len(self.prs.slide_layouts) - 1)
    
    def _get_paragraph_style(self, placeholder) -> Dict[str, Any]:
        """获取占位符的段落样式"""
        style = {}
        try:
            if hasattr(placeholder, 'text_frame') and placeholder.text_frame:
                tf = placeholder.text_frame
                if tf.paragraphs:
                    p = tf.paragraphs[0]
                    if p.font.name:
                        style['font_name'] = p.font.name
                    if p.font.size:
                        style['font_size'] = p.font.size
                    if p.font.bold is not None:
                        style['font_bold'] = p.font.bold
                    if p.font.italic is not None:
                        style['font_italic'] = p.font.italic
                    try:
                        if p.font.color and p.font.color.rgb:
                            style['font_color'] = p.font.color.rgb
                    except:
                        pass
                    if p.alignment is not None:
                        style['alignment'] = p.alignment
        except Exception as e:
            logger.debug(f"[PPTXBuilder] 获取样式失败: {e}")
        return style
    
    def _apply_paragraph_style(self, paragraph, style: Dict[str, Any]):
        """应用段落样式"""
        try:
            if 'font_name' in style:
                paragraph.font.name = style['font_name']
            if 'font_size' in style:
                paragraph.font.size = style['font_size']
            if 'font_bold' in style:
                paragraph.font.bold = style['font_bold']
            if 'font_italic' in style:
                paragraph.font.italic = style['font_italic']
            if 'font_color' in style:
                try:
                    paragraph.font.color.rgb = style['font_color']
                except:
                    pass
            if 'alignment' in style:
                paragraph.alignment = style['alignment']
        except Exception as e:
            logger.debug(f"[PPTXBuilder] 应用样式失败: {e}")
    
    def add_slide_from_data(self, slide_data: Dict):
        """根据数据添加幻灯片，保留样式"""
        slide_type = slide_data.get('type', 'content').lower()
        layout_idx = self._find_suitable_layout(slide_type)
        slide_layout = self.prs.slide_layouts[layout_idx]
        
        slide = self.prs.slides.add_slide(slide_layout)
        
        title_text = slide_data.get('title', '')
        if slide.shapes.title:
            title_style = self._get_paragraph_style(slide.shapes.title)
            slide.shapes.title.text = title_text
            if slide.shapes.title.text_frame and slide.shapes.title.text_frame.paragraphs:
                self._apply_paragraph_style(slide.shapes.title.text_frame.paragraphs[0], title_style)
        
        content_text = slide_data.get('content', '')
        bullets = slide_data.get('bullets', [])
        
        body_ph = None
        body_style = {}
        
        for shape in slide.placeholders:
            ph_type = self._get_placeholder_type(shape)
            if ph_type == 'body':
                body_ph = shape
                body_style = self._get_paragraph_style(shape)
                break
        
        if not body_ph:
            for shape in slide.placeholders:
                if shape != slide.shapes.title and hasattr(shape, 'text_frame'):
                    body_ph = shape
                    body_style = self._get_paragraph_style(shape)
                    break
        
        if body_ph:
            self._fill_content(body_ph, content_text, bullets, body_style)
        
        image_prompt = slide_data.get('imagePrompt')
        if image_prompt:
            logger.debug(f"[PPTXBuilder] 幻灯片 '{title_text}' 需要图片: {image_prompt}")
    
    def _fill_content(self, placeholder, content_text: str, bullets: List[str], style: Dict[str, Any]):
        """填充内容并保留样式"""
        try:
            tf = placeholder.text_frame
            
            if bullets and isinstance(bullets, list):
                tf.clear()
                for i, bullet in enumerate(bullets):
                    if i == 0:
                        p = tf.paragraphs[0]
                    else:
                        p = tf.add_paragraph()
                    p.text = str(bullet)
                    p.level = 0
                    self._apply_paragraph_style(p, style)
            elif content_text:
                tf.clear()
                p = tf.paragraphs[0]
                p.text = str(content_text)
                self._apply_paragraph_style(p, style)
                
        except Exception as e:
            logger.warning(f"[PPTXBuilder] 填充内容失败: {e}")
            if bullets:
                placeholder.text = "\n".join(bullets)
            else:
                placeholder.text = str(content_text)
    
    def build_stream(self) -> io.BytesIO:
        """生成内存流"""
        stream = io.BytesIO()
        self.prs.save(stream)
        stream.seek(0)
        logger.info(f"[PPTXBuilder] 生成完成，文件大小: {len(stream.getvalue())} 字节")
        return stream


def create_pptx_from_ai_json(slides_data: List[Dict], template_id: Optional[str] = None) -> io.BytesIO:
    """
    高层封装：将JSON数据转为PPTX流
    
    Args:
        slides_data: 幻灯片数据列表
        template_id: 模板ID（可选）
        
    Returns:
        PPTX文件的字节流
    """
    template_path = None
    if template_id:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        template_path = os.path.join(base_dir, 'data', 'templates', f"{template_id}.pptx")
        if not os.path.exists(template_path):
            logger.warning(f"[PPTXBuilder] 模板 {template_id} 不存在")
            template_path = None

    builder = PPTXBuilder(template_path)
    for slide_data in slides_data:
        builder.add_slide_from_data(slide_data)
        
    return builder.build_stream()
