"""
PPT 幻灯片渲染器 - 将 PPT 转换为可预览的图片
支持多种渲染方式：
1. PowerPoint COM 接口（Windows + PowerPoint）
2. LibreOffice 命令行
3. python-pptx 文本提取（纯 Python，无需外部依赖）
"""
import os
import io
import base64
import logging
import tempfile
import shutil
from typing import List, Dict, Any, Optional
from PIL import Image, ImageDraw, ImageFont

logger = logging.getLogger(__name__)

COMTYPES_AVAILABLE = False
try:
    import comtypes.client
    COMTYPES_AVAILABLE = True
    logger.info("comtypes 已加载，PowerPoint COM 渲染可用")
except ImportError:
    logger.info("comtypes 未安装，PowerPoint COM 渲染不可用")

PYTHON_PPTX_AVAILABLE = False
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    PYTHON_PPTX_AVAILABLE = True
    logger.info("python-pptx 已加载，文本预览可用")
except ImportError:
    logger.info("python-pptx 未安装，文本预览不可用")


class SlideRenderer:
    """PPT幻灯片渲染器"""
    
    SLIDE_WIDTH = 1280
    SLIDE_HEIGHT = 720
    
    @staticmethod
    def render_pptx_to_images(file_path: str, max_slides: int = 50) -> List[Dict[str, Any]]:
        """
        将 PPT 文件渲染为图片列表
        
        Args:
            file_path: PPT 文件路径
            max_slides: 最大渲染幻灯片数量
            
        Returns:
            List of slide data with image base64 and info
        """
        abs_path = os.path.abspath(file_path)
        
        logger.info(f"[SlideRenderer] 开始渲染: {abs_path}")
        logger.info(f"[SlideRenderer] COM可用: {COMTYPES_AVAILABLE}, PPTX可用: {PYTHON_PPTX_AVAILABLE}")
        
        if not os.path.exists(abs_path):
            logger.error(f"PPT 文件不存在: {abs_path}")
            raise FileNotFoundError(f"PPT 文件不存在: {abs_path}")
        
        file_size = os.path.getsize(abs_path)
        logger.info(f"[SlideRenderer] 文件大小: {file_size} 字节")
        
        if file_size < 1000:
            logger.warning(f"[SlideRenderer] 文件过小，可能已损坏")
        
        if COMTYPES_AVAILABLE:
            try:
                logger.info("[SlideRenderer] 尝试使用 PowerPoint COM 渲染...")
                slides_data = SlideRenderer._render_with_com(abs_path, max_slides)
                logger.info(f"[SlideRenderer] COM 渲染成功: {len(slides_data)} 页")
                return slides_data
            except Exception as e:
                logger.warning(f"[SlideRenderer] COM 渲染失败: {type(e).__name__}: {e}, 尝试备用方案...")
        
        logger.info("[SlideRenderer] 使用备用渲染方案...")
        return SlideRenderer._render_fallback(abs_path, max_slides)
    
    @staticmethod
    def _render_with_com(file_path: str, max_slides: int) -> List[Dict[str, Any]]:
        """使用 PowerPoint COM 接口渲染"""
        import comtypes.client
        
        temp_dir = tempfile.mkdtemp()
        slides_data = []
        
        try:
            powerpoint = comtypes.client.CreateObject("PowerPoint.Application")
            powerpoint.Visible = True
            
            presentation = powerpoint.Presentations.Open(file_path, WithWindow=False)
            
            total_slides = min(presentation.Slides.Count, max_slides)
            
            for i in range(1, total_slides + 1):
                slide = presentation.Slides.Item(i)
                slide_path = os.path.join(temp_dir, f"slide_{i}.png")
                
                slide.Export(slide_path, "PNG", SlideRenderer.SLIDE_WIDTH, SlideRenderer.SLIDE_HEIGHT)
                
                with open(slide_path, 'rb') as f:
                    image_data = f.read()
                    image_base64 = base64.b64encode(image_data).decode('utf-8')
                
                slide_info = {
                    'slide_num': i,
                    'title': f'第 {i} 页',
                    'content_preview': '',
                    'image': f"data:image/png;base64,{image_base64}"
                }
                slides_data.append(slide_info)
            
            presentation.Close()
            powerpoint.Quit()
            
        except Exception as e:
            logger.error(f"COM 渲染异常: {e}")
            raise
        finally:
            try:
                shutil.rmtree(temp_dir, ignore_errors=True)
            except:
                pass
        
        return slides_data
    
    @staticmethod
    def _render_fallback(file_path: str, max_slides: int) -> List[Dict[str, Any]]:
        """备用渲染方案 - 依次尝试 LibreOffice、python-pptx"""
        temp_dir = tempfile.mkdtemp()
        slides_data = []
        
        try:
            logger.info("[SlideRenderer] 尝试 LibreOffice 渲染...")
            result = SlideRenderer._convert_with_libreoffice(file_path, temp_dir)
            
            if result:
                image_files = sorted([f for f in os.listdir(temp_dir) if f.endswith('.png')])
                logger.info(f"[SlideRenderer] LibreOffice 生成了 {len(image_files)} 个图片文件")
                
                for i, img_file in enumerate(image_files[:max_slides]):
                    img_path = os.path.join(temp_dir, img_file)
                    
                    with Image.open(img_path) as img:
                        img = img.resize((SlideRenderer.SLIDE_WIDTH, SlideRenderer.SLIDE_HEIGHT), Image.Resampling.LANCZOS)
                        buffer = io.BytesIO()
                        img.save(buffer, format='PNG')
                        image_base64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
                    
                    slide_info = {
                        'slide_num': i + 1,
                        'title': f'第 {i + 1} 页',
                        'content_preview': '',
                        'image': f"data:image/png;base64,{image_base64}"
                    }
                    slides_data.append(slide_info)
                
                if slides_data:
                    logger.info(f"[SlideRenderer] LibreOffice 渲染完成: {len(slides_data)} 页")
                    return slides_data
            else:
                logger.info("[SlideRenderer] LibreOffice 渲染失败或未安装")
            
            if PYTHON_PPTX_AVAILABLE:
                logger.info("[SlideRenderer] 尝试 python-pptx 文本预览...")
                try:
                    slides_data = SlideRenderer._render_with_pptx(file_path, max_slides)
                    if slides_data and len(slides_data) > 0:
                        logger.info(f"[SlideRenderer] python-pptx 文本预览完成: {len(slides_data)} 页")
                        return slides_data
                    else:
                        logger.warning("[SlideRenderer] python-pptx 返回空结果")
                except Exception as pptx_err:
                    logger.error(f"[SlideRenderer] python-pptx 渲染异常: {type(pptx_err).__name__}: {pptx_err}")
            else:
                logger.warning("[SlideRenderer] python-pptx 不可用")
            
            logger.warning("[SlideRenderer] 所有渲染方案失败，尝试读取 PPT 页数生成占位符")
            slide_count = SlideRenderer._get_pptx_slide_count(file_path)
            logger.info(f"[SlideRenderer] PPT 页数: {slide_count}")
            if slide_count > 0:
                return SlideRenderer._create_placeholder_slides(slide_count)
            
            return SlideRenderer._create_placeholder_slides(1)
            
        except Exception as e:
            logger.error(f"[SlideRenderer] 备用渲染失败: {type(e).__name__}: {e}", exc_info=True)
            return SlideRenderer._create_placeholder_slides(1)
        finally:
            try:
                shutil.rmtree(temp_dir, ignore_errors=True)
            except:
                pass
    
    @staticmethod
    def _get_pptx_slide_count(file_path: str) -> int:
        """获取 PPT 文件的幻灯片数量"""
        if not PYTHON_PPTX_AVAILABLE:
            return 0
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            return len(prs.slides)
        except Exception as e:
            logger.error(f"获取 PPT 页数失败: {e}")
            return 0
    
    @staticmethod
    def _convert_with_libreoffice(file_path: str, output_dir: str) -> bool:
        """使用 LibreOffice 转换 PPT 为图片"""
        import subprocess
        
        libreoffice_paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            "soffice",
            "libreoffice"
        ]
        
        for lo_path in libreoffice_paths:
            try:
                cmd = [
                    lo_path,
                    "--headless",
                    "--convert-to", "png",
                    "--outdir", output_dir,
                    file_path
                ]
                
                result = subprocess.run(cmd, capture_output=True, timeout=60)
                
                if result.returncode == 0:
                    return True
            except FileNotFoundError:
                continue
            except subprocess.TimeoutExpired:
                continue
            except Exception as e:
                logger.debug(f"LibreOffice 路径 {lo_path} 失败: {e}")
                continue
        
        return False
    
    @staticmethod
    def _render_with_pptx(file_path: str, max_slides: int) -> List[Dict[str, Any]]:
        """使用 python-pptx 提取内容和样式生成预览图（优化版）"""
        from pptx import Presentation
        from pptx.util import Pt, Inches, Emu
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.dml.color import RGBColor
        
        try:
            logger.info(f"[SlideRenderer._render_with_pptx] 加载 PPT: {file_path}")
            prs = Presentation(file_path)
            logger.info(f"[SlideRenderer._render_with_pptx] PPT 页数: {len(prs.slides)}")
            
            slide_width_emu = prs.slide_width
            slide_height_emu = prs.slide_height
            
            slides_data = []
            
            total_slides = min(len(prs.slides), max_slides)
            for i in range(total_slides):
                try:
                    slide = prs.slides[i]
                    slide_info = SlideRenderer._render_single_slide(
                        slide, i + 1, slide_width_emu, slide_height_emu
                    )
                    slides_data.append(slide_info)
                except Exception as e:
                    logger.warning(f"[SlideRenderer._render_with_pptx] 渲染第 {i+1} 页失败: {e}")
                    continue
            
            logger.info(f"[SlideRenderer._render_with_pptx] 完成，共 {len(slides_data)} 页")
            return slides_data
            
        except Exception as e:
            logger.error(f"[SlideRenderer._render_with_pptx] 失败: {type(e).__name__}: {e}", exc_info=True)
            return []
    
    @staticmethod
    def _render_single_slide(slide, slide_num: int, slide_width_emu, slide_height_emu) -> Dict[str, Any]:
        """渲染单个幻灯片为图片"""
        scale_x = SlideRenderer.SLIDE_WIDTH / slide_width_emu
        scale_y = SlideRenderer.SLIDE_HEIGHT / slide_height_emu
        
        bg_color = SlideRenderer._get_slide_background_color(slide)
        
        img = Image.new('RGB', (SlideRenderer.SLIDE_WIDTH, SlideRenderer.SLIDE_HEIGHT), bg_color)
        draw = ImageDraw.Draw(img)
        
        title_text = ""
        content_texts = []
        
        shapes_data = []
        for shape in slide.shapes:
            shape_data = SlideRenderer._extract_shape_data(shape, scale_x, scale_y)
            if shape_data:
                shapes_data.append(shape_data)
        
        shapes_data.sort(key=lambda x: (x.get('top', 0), x.get('left', 0)))
        
        for shape_data in shapes_data:
            SlideRenderer._draw_shape(draw, shape_data, img)
            
            if shape_data.get('is_title') and shape_data.get('text'):
                title_text = shape_data['text']
            elif shape_data.get('text'):
                content_texts.append(shape_data['text'])
        
        buffer = io.BytesIO()
        img.save(buffer, format='PNG')
        image_base64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
        
        content_preview = "\n".join(content_texts[:5])[:200]
        
        return {
            'slide_num': slide_num,
            'title': title_text or f'第 {slide_num} 页',
            'content_preview': content_preview,
            'image': f"data:image/png;base64,{image_base64}"
        }
    
    @staticmethod
    def _get_slide_background_color(slide) -> str:
        """获取幻灯片背景颜色"""
        try:
            if hasattr(slide, 'background') and slide.background:
                fill = slide.background.fill
                if fill.type:
                    if hasattr(fill, 'fore_color') and fill.fore_color:
                        if hasattr(fill.fore_color, 'rgb') and fill.fore_color.rgb:
                            rgb = fill.fore_color.rgb
                            return f'#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}'
        except Exception:
            pass
        return '#FFFFFF'
    
    @staticmethod
    def _extract_shape_data(shape, scale_x: float, scale_y: float) -> Optional[Dict[str, Any]]:
        """提取形状数据"""
        try:
            data = {
                'name': shape.name,
                'left': int(shape.left * scale_x) if hasattr(shape, 'left') and shape.left else 0,
                'top': int(shape.top * scale_y) if hasattr(shape, 'top') and shape.top else 0,
                'width': int(shape.width * scale_x) if hasattr(shape, 'width') and shape.width else 100,
                'height': int(shape.height * scale_y) if hasattr(shape, 'height') and shape.height else 50,
                'shape_type': str(shape.shape_type) if hasattr(shape, 'shape_type') else 'UNKNOWN',
            }
            
            data['is_title'] = 'title' in shape.name.lower() if shape.name else False
            
            if hasattr(shape, 'fill') and shape.fill.type:
                try:
                    if hasattr(shape.fill, 'fore_color') and shape.fill.fore_color:
                        if hasattr(shape.fill.fore_color, 'rgb') and shape.fill.fore_color.rgb:
                            rgb = shape.fill.fore_color.rgb
                            data['fill_color'] = f'#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}'
                except Exception:
                    pass
            
            if hasattr(shape, 'text_frame') and shape.text_frame:
                text_info = SlideRenderer._extract_text_info(shape.text_frame)
                data.update(text_info)
            
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                data['is_picture'] = True
                try:
                    if hasattr(shape, 'image') and shape.image:
                        image_bytes = shape.image.blob
                        img = Image.open(io.BytesIO(image_bytes))
                        data['image_data'] = img
                except Exception:
                    pass
            
            return data
        except Exception as e:
            logger.debug(f"提取形状数据失败: {e}")
            return None
    
    @staticmethod
    def _extract_text_info(text_frame) -> Dict[str, Any]:
        """提取文本信息"""
        result = {
            'text': '',
            'font_size': 18,
            'font_color': '#333333',
            'font_name': 'Arial',
            'bold': False,
            'alignment': 'left'
        }
        
        try:
            paragraphs_text = []
            for paragraph in text_frame.paragraphs:
                para_text = ""
                for run in paragraph.runs:
                    para_text += run.text
                
                if paragraph.runs:
                    run = paragraph.runs[0]
                    if hasattr(run.font, 'size') and run.font.size:
                        result['font_size'] = int(run.font.size.pt)
                    if hasattr(run.font, 'name') and run.font.name:
                        result['font_name'] = run.font.name
                    if hasattr(run.font, 'bold'):
                        result['bold'] = run.font.bold
                    if hasattr(run.font, 'color') and run.font.color:
                        if hasattr(run.font.color, 'rgb') and run.font.color.rgb:
                            rgb = run.font.color.rgb
                            result['font_color'] = f'#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}'
                
                paragraphs_text.append(para_text)
            
            result['text'] = "\n".join(paragraphs_text)
            
        except Exception as e:
            logger.debug(f"提取文本信息失败: {e}")
        
        return result
    
    @staticmethod
    def _draw_shape(draw: ImageDraw.ImageDraw, shape_data: Dict[str, Any], img: Image.Image):
        """绘制形状到图片"""
        try:
            left = max(0, shape_data.get('left', 0))
            top = max(0, shape_data.get('top', 0))
            width = min(shape_data.get('width', 100), SlideRenderer.SLIDE_WIDTH - left)
            height = min(shape_data.get('height', 50), SlideRenderer.SLIDE_HEIGHT - top)
            right = left + width
            bottom = top + height
            
            if shape_data.get('is_picture') and shape_data.get('image_data'):
                try:
                    pic_img = shape_data['image_data']
                    pic_img = pic_img.resize((width, height), Image.Resampling.LANCZOS)
                    img.paste(pic_img, (left, top))
                except Exception:
                    draw.rectangle([left, top, right, bottom], fill='#E0E0E0', outline='#CCCCCC')
                return
            
            fill_color = shape_data.get('fill_color')
            if fill_color:
                draw.rectangle([left, top, right, bottom], fill=fill_color)
            
            text = shape_data.get('text', '')
            if text:
                font_size = min(shape_data.get('font_size', 18), 48)
                font_color = shape_data.get('font_color', '#333333')
                font_name = shape_data.get('font_name', 'Arial')
                bold = shape_data.get('bold', False)
                
                try:
                    font_paths = [
                        f"C:\\Windows\\Fonts\\{font_name}.ttf",
                        f"C:\\Windows\\Fonts\\simhei.ttf",
                        "arial.ttf",
                    ]
                    font = None
                    for fp in font_paths:
                        try:
                            font = ImageFont.truetype(fp, font_size)
                            break
                        except Exception:
                            continue
                    if not font:
                        font = ImageFont.load_default()
                except Exception:
                    font = ImageFont.load_default()
                
                lines = text.split('\n')
                y_offset = top + 5
                max_lines = min(len(lines), height // (font_size + 4))
                
                for line in lines[:max_lines]:
                    if y_offset + font_size > bottom - 5:
                        break
                    
                    display_text = line[:int(width / (font_size * 0.5))]
                    
                    draw.text((left + 8, y_offset), display_text, fill=font_color, font=font)
                    y_offset += font_size + 4
            
        except Exception as e:
            logger.debug(f"绘制形状失败: {e}")
    
    @staticmethod
    def _create_text_slide_image(slide_num: int, text: str) -> Image.Image:
        """根据文本内容创建幻灯片预览图"""
        img = Image.new('RGB', (SlideRenderer.SLIDE_WIDTH, SlideRenderer.SLIDE_HEIGHT), '#ffffff')
        draw = ImageDraw.Draw(img)
        
        try:
            title_font = ImageFont.truetype("arial.ttf", 36)
            content_font = ImageFont.truetype("arial.ttf", 18)
        except:
            try:
                title_font = ImageFont.truetype("C:\\Windows\\Fonts\\arial.ttf", 36)
                content_font = ImageFont.truetype("C:\\Windows\\Fonts\\arial.ttf", 18)
            except:
                title_font = ImageFont.load_default()
                content_font = ImageFont.load_default()
        
        draw.rectangle([0, 0, SlideRenderer.SLIDE_WIDTH, 60], fill='#4472C4')
        title = f"第 {slide_num} 页"
        bbox = draw.textbbox((0, 0), title, font=title_font)
        title_width = bbox[2] - bbox[0]
        draw.text(((SlideRenderer.SLIDE_WIDTH - title_width) // 2, 12), title, fill='#ffffff', font=title_font)
        
        y_offset = 80
        lines = text.split('\n')
        for line in lines:
            if y_offset > SlideRenderer.SLIDE_HEIGHT - 40:
                break
            
            words = line
            max_chars = 80
            while len(words) > max_chars:
                draw.text((40, y_offset), words[:max_chars], fill='#333333', font=content_font)
                words = words[max_chars:]
                y_offset += 28
            if words:
                draw.text((40, y_offset), words, fill='#333333', font=content_font)
                y_offset += 28
        
        return img
    
    @staticmethod
    def _create_placeholder_slides(count: int) -> List[Dict[str, Any]]:
        """创建占位符幻灯片"""
        img = Image.new('RGB', (SlideRenderer.SLIDE_WIDTH, SlideRenderer.SLIDE_HEIGHT), '#f0f0f0')
        draw = ImageDraw.Draw(img)
        
        try:
            font = ImageFont.truetype("arial.ttf", 48)
        except:
            font = ImageFont.load_default()
        
        slides_data = []
        
        for i in range(1, min(count + 1, 51)):
            slide_img = img.copy()
            draw = ImageDraw.Draw(slide_img)
            
            text = f"第 {i} 页"
            bbox = draw.textbbox((0, 0), text, font=font)
            text_width = bbox[2] - bbox[0]
            text_height = bbox[3] - bbox[1]
            
            x = (SlideRenderer.SLIDE_WIDTH - text_width) // 2
            y = (SlideRenderer.SLIDE_HEIGHT - text_height) // 2
            
            draw.text((x, y), text, fill='#666666', font=font)
            
            buffer = io.BytesIO()
            slide_img.save(buffer, format='PNG')
            image_base64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
            
            slide_info = {
                'slide_num': i,
                'title': f'第 {i} 页',
                'content_preview': '预览不可用，请下载后查看',
                'image': f"data:image/png;base64,{image_base64}"
            }
            slides_data.append(slide_info)
        
        return slides_data
