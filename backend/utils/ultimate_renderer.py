"""
终极 PPT 渲染引擎 v3 - 真正的“填字”版本
==========================================

核心原理（v3 重写）：
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
❌ 旧版（v1/v2）问题：
   TemplateStyleCloner.clone_style_to_new_presentation()
   → 把模板完整复制到新 Presentation（包含原始示例页）
   → 用 XML 操作尝试删除原始页（不可靠！python-pptx 不保证清理干净）
   → 结果：生成的 PPT 里混着模板的原始内容（"粘贴模板"效果）

✅ v3 方案：Add-then-Clean（先添加后删除）
   1. 直接打开模板 Presentation（保留所有样式：母版/版式/颜色/字体）
   2. 记录原始模板有多少页
   3. 用模板的版式逐个 add_slide() → 创建干净的空幻灯片（只有占位符框）
   4. 向占位符填入用户的文字（标题、正文、副标题）
   5. 所有用户页面添加完毕后，删除原始模板页
   6. 保存输出

为什么这能工作：
- add_slide(layout) 永远创建空白幻灯片（只有空的占位符位置）
- 填字操作只修改占位符文本，不碰任何格式属性
- 删除原始页是在最后一步，不存在中间状态不一致的问题
"""

import os
import io
import logging
from typing import List, Dict, Any, Optional, Tuple

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

logger = logging.getLogger(__name__)

LAYOUT_KEYWORD_MAP = {
    'cover': ['title slide', '封面', '标题页', 'title only', 'cover', 'blank'],
    'toc': ['toc', 'agenda', '目录', 'section header'],
    'content': ['title and content', 'content', 'two content', 'bullet', 'text'],
    'summary': ['summary', '总结', 'conclusion', 'key points'],
    'ending': ['blank', 'ending', 'closing', 'thank', '致谢', 'thank you'],
}


class UltimateRenderer:
    """
    终极 PPT 渲染器 v3 — 在模板版式中填入用户内容
    
    工作流程：
    1. 打开模板文件（继承全部视觉样式）
    2. 逐页创建新幻灯片并填入用户文字
    3. 删除所有原始模板页
    4. 输出只含用户内容的 PPTX
    """

    def __init__(self, template_path: str):
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"模板文件不存在: {template_path}")

        self.template_path = template_path
        self.prs: Optional[Presentation] = None
        self.original_slide_count: int = 0
        self._layout_info: List[Dict[str, Any]] = []

        self._initialize()

    def _initialize(self):
        """直接打开模板，记录原始状态"""
        logger.info("[UltimateRenderer-v3] 初始化...")
        logger.info(f"[UltimateRenderer-v3]   模板路径: {self.template_path}")

        self.prs = Presentation(self.template_path)
        self.original_slide_count = len(self.prs.slides)

        self._analyze_layouts()

        logger.info(f"[UltimateRenderer-v3] 初始化完成:")
        logger.info(f"   - 版式数量: {len(self.prs.slide_layouts)}")
        logger.info(f"   - 母版数量: {len(self.prs.slide_masters)}")
        logger.info(f"   - 原始模板页数: {self.original_slide_count}")

    def _analyze_layouts(self):
        """分析每个版式的特征，用于智能匹配"""
        self._layout_info = []

        for idx, layout in enumerate(self.prs.slide_layouts):
            info = {
                'index': idx,
                'name': layout.name or '',
                'has_title': False,
                'has_body': False,
                'has_subtitle': False,
                'placeholder_map': {},
            }

            for shape in layout.placeholders:
                try:
                    ph_type = str(shape.placeholder_format.type)
                    ph_idx = shape.placeholder_format.idx
                    ph_name = shape.name or ''

                    info['placeholder_map'][ph_idx] = {
                        'type': ph_type,
                        'name': ph_name,
                    }

                    if ph_type in (str(PP_PLACEHOLDER.TITLE), str(PP_PLACEHOLDER.CENTER_TITLE),
                                   str(PP_PLACEHOLDER.VERTICAL_TITLE)):
                        info['has_title'] = True
                    elif ph_type == str(PP_PLACEHOLDER.SUBTITLE):
                        info['has_subtitle'] = True
                    elif ph_type in (str(PP_PLACEHOLDER.BODY), str(PP_PLACEHOLDER.OBJECT)):
                        info['has_body'] = True
                except Exception:
                    pass

            self._layout_info.append(info)
            logger.debug(f"[UltimateRenderer-v3] 版式[{idx}] {info['name']}: "
                        f"title={info['has_title']} body={info['has_body']} "
                        f"placeholders={len(info['placeholder_map'])}")

    def _select_best_layout(self, page_type: str) -> Tuple[Any, Dict[str, Any]]:
        """根据页面类型选择最合适的版式 v3.1 - 修复版式选择逻辑"""
        keywords = LAYOUT_KEYWORD_MAP.get(page_type.lower(), LAYOUT_KEYWORD_MAP['content'])

        needs_body = page_type.lower() in ('content', 'toc', 'summary')
        
        candidates = []
        
        for idx, layout in enumerate(self.prs.slide_layouts):
            info = self._layout_info[idx] if idx < len(self._layout_info) else {}
            name_lower = (info.get('name') or '').lower()
            score = 0

            for kw in keywords:
                if kw.lower() in name_lower:
                    score += 10

            if page_type.lower() == 'cover' and info.get('has_title') and not info.get('has_body'):
                score += 5
            elif needs_body and info.get('has_title') and info.get('has_body'):
                score += 5
            elif page_type.lower() == 'ending' and info.get('has_title') and not info.get('has_body'):
                score += 5
            
            if needs_body and info.get('has_body'):
                score += 10

            candidates.append((score, layout, info))

        candidates.sort(key=lambda x: x[0], reverse=True)
        
        best_score, best_layout, best_info = candidates[0] if candidates else (0, None, None)

        if needs_body and best_layout is not None and not best_info.get('has_body'):
            logger.warning(
                f"[UltimateRenderer-v3.1] 最佳版式缺少正文占位符 | "
                f"type={page_type} | chosen='{best_layout.name}' | "
                f"尝试寻找替代版式..."
            )
            
            body_candidates = [(s, l, i) for s, l, i in candidates if i.get('has_body')]
            if body_candidates:
                preferred = None
                for s, l, i in body_candidates:
                    if i.get('has_title'):
                        preferred = (s, l, i)
                        break
                
                if not preferred:
                    preferred = body_candidates[0]
                
                best_score, best_layout, best_info = preferred
                logger.info(
                    f"[UltimateRenderer-v3.1] 切换到带正文占位符的版式: '{best_layout.name}'"
                )

        if best_layout is None:
            for idx, layout in enumerate(self.prs.slide_layouts):
                info = self._layout_info[idx]
                if info.get('has_title') and info.get('has_body'):
                    return layout, info

            if self.prs.slide_layouts:
                return self.prs.slide_layouts[0], self._layout_info[0]

            raise RuntimeError("模板没有任何可用版式")

        logger.debug(f"[UltimateRenderer-v3.1] 选中的版式: {best_layout.name} (score={best_score})")
        return best_layout, best_info

    def render(self, slides_data: List[Dict[str, Any]]) -> io.BytesIO:
        """
        渲染幻灯片列表
        
        Args:
            slides_data: 幻灯片数据列表
                [
                    {
                        'title': '标题',
                        'subtitle': '副标题',
                        'content': ['要点1', '要点2'],
                        'page_type': 'cover'
                    },
                    ...
                ]
                
        Returns:
            PPTX 字节流
        """
        logger.info(f"[UltimateRenderer-v3] 开始渲染 {len(slides_data)} 页...")

        for idx, slide_dict in enumerate(slides_data):
            page_type = slide_dict.get('page_type', 'content')
            title = slide_dict.get('title', f'第{idx+1}页')
            content = slide_dict.get('content', [])
            subtitle = slide_dict.get('subtitle', '')

            layout, layout_info = self._select_best_layout(page_type)

            slide = self.prs.slides.add_slide(layout)

            self._fill_slide(slide, title, subtitle, content, layout_info)

            logger.info(f"[UltimateRenderer-v3] 第 {idx+1}/{len(slides_data)} 页完成: "
                       f"{title[:30]}... (版式: {layout.name})")

        self._remove_original_slides()

        result = io.BytesIO()
        self.prs.save(result)
        result.seek(0)

        file_size = len(result.getvalue())
        total_slides = len(self.prs.slides)
        logger.info(f"[UltimateRenderer-v3] 渲染完成!")
        logger.info(f"   文件大小: {file_size / 1024:.1f} KB")
        logger.info(f"   最终幻灯片数: {total_slides} (用户{len(slides_data)}页)")

        return result

    def _fill_slide(self, slide, title: str, subtitle: str,
                     content: List[str], layout_info: Dict[str, Any]):
        """
        向幻灯片的占位符中填入用户内容
        
        核心原则：只修改 text，不修改任何格式属性
        这样字体、颜色、大小、对齐方式全部继承自模板版式
        """
        ph_map = layout_info.get('placeholder_map', {})

        for shape in slide.placeholders:
            try:
                ph_idx = shape.placeholder_format.idx
                ph_info = ph_map.get(ph_idx, {})
                ph_type_str = ph_info.get('type', '')

                if not hasattr(shape, 'text_frame'):
                    continue

                tf = shape.text_frame

                if ph_type_str in (str(PP_PLACEHOLDER.TITLE), str(PP_PLACEHOLDER.CENTER_TITLE),
                                    str(PP_PLACEHOLDER.VERTICAL_TITLE)):

                    if title:
                        self._set_text_to_placeholder(tf, title)

                elif ph_type_str == str(PP_PLACEHOLDER.SUBTITLE):

                    if subtitle:
                        self._set_text_to_placeholder(tf, subtitle)

                elif ph_type_str in (str(PP_PLACEHOLDER.BODY), str(PP_PLACEHOLDER.OBJECT)):

                    if content:
                        self._set_body_to_placeholder(tf, content)

            except Exception as e:
                logger.debug(f"[UltimateRenderer-v3] 填充占位符失败(idx={shape.placeholder_format.idx}): {e}")

    def _set_text_to_placeholder(self, tf, text: str):
        """向占位符设置纯文本（保持原格式）"""
        try:
            if not tf.paragraphs:
                para = tf.add_paragraph()
            else:
                para = tf.paragraphs[0]

            if para.runs:
                para.runs[0].text = text
                for run in list(para.runs)[1:]:
                    run._element.getparent().remove(run._element)
            else:
                run = para.add_run()
                run.text = text

            for p in list(tf.paragraphs)[1:]:
                p._element.getparent().remove(p._element)

        except Exception as e:
            logger.debug(f"[UltimateRenderer-v3] 设置文本失败: {e}")
            try:
                tf.text = text
            except Exception:
                pass

    def _set_body_to_placeholder(self, tf, lines: List[str]):
        """向正文占位符设置多行内容（保持项目符号等列表格式）"""
        try:
            first_para_fmt = None
            if tf.paragraphs:
                p0 = tf.paragraphs[0]
                first_para_fmt = {
                    'level': getattr(p0, 'level', 0),
                    'alignment': getattr(p0, 'alignment', None),
                }

            for p in list(tf.paragraphs):
                p._element.getparent().remove(p._element)

            for i, line in enumerate(lines):
                para = tf.add_paragraph()

                if first_para_fmt:
                    para.level = first_para_fmt.get('level', 0)
                    if first_para_fmt.get('alignment'):
                        para.alignment = first_para_fmt['alignment']

                run = para.add_run()
                run.text = line

        except Exception as e:
            logger.debug(f"[UltimateRenderer-v3] 设置正文失败: {e}")
            try:
                tf.text = '\n'.join(lines)
            except Exception:
                pass

    def _remove_original_slides(self):
        """
        删除所有原始模板幻灯片（在用户页面已安全添加后执行）
        
        使用 python-pptx 内部 API (prs.slides._sldIdLst) 直接操作
        """
        removed = 0

        for _ in range(self.original_slide_count):
            if len(self.prs.slides) <= 0:
                break

            try:
                sld_id_lst = self.prs.slides._sldIdLst
                rId = sld_id_lst[0].rId
                self.prs.part.drop_rel(rId)
                del sld_id_lst[0]
                removed += 1
            except Exception as e:
                logger.warning(f"[UltimateRenderer-v3] 删除原始页出错: {e}")
                break

        logger.info(f"[UltimateRenderer-v3] 已删除 {removed} 张原始模板页")


def render_ultimate_ppt(
    template_path: str,
    slides_data: List[Dict[str, Any]],
    output_path: str = None
) -> io.BytesIO:
    """
    高层接口：使用终极渲染器生成 PPT
    
    Args:
        template_path: 模板文件路径
        slides_data: 幻灯片数据列表
        output_path: 可选的输出路径
        
    Returns:
        PPTX 字节流
    """
    logger.info("=" * 60)
    logger.info("[UltimatePPT-v3] 开始生成...")
    logger.info(f"   模板: {template_path}")
    logger.info(f"   幻灯片数: {len(slides_data)}")

    renderer = UltimateRenderer(template_path)
    result = renderer.render(slides_data)

    if output_path:
        with open(output_path, 'wb') as f:
            f.write(result.getvalue())
        logger.info(f"[UltimatePPT-v3] 已保存到: {output_path}")
        result.seek(0)

    logger.info("=" * 60)
    return result


if __name__ == "__main__":
    import sys

    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
    )

    print("=" * 60)
    print("终极 PPT 渲染器 v3 - 测试程序")
    print("策略: Add-then-Clean（先添加用户页，再删除原始模板页）")
    print("=" * 60)

    if len(sys.argv) < 2:
        print("\n用法:")
        print("  python ultimate_renderer.py <模板路径>")
        sys.exit(1)

    template_path = sys.argv[1]

    test_slides = [
        {
            'title': '教学演示文稿',
            'subtitle': '使用模板样式自动排版',
            'page_type': 'cover'
        },
        {
            'title': '目录',
            'content': [
                '一、课程目标与要求',
                '二、核心知识点讲解',
                '三、课堂互动练习',
                '四、课后作业布置'
            ],
            'page_type': 'toc'
        },
        {
            'title': '第一章：核心知识点',
            'content': [
                '本章节将介绍以下重要概念：',
                '• 概念一：基础定义与应用场景',
                '• 概念二：进阶技巧与实践方法',
                '• 概念三：常见误区与避坑指南'
            ],
            'page_type': 'content'
        },
        {
            'title': '课堂练习',
            'content': [
                '请同学们完成以下练习题：',
                '',
                '1. 根据所学内容，分析案例A的解决方案',
                '2. 小组讨论：如何优化现有流程？',
                '3. 分享你的思考结果'
            ],
            'page_type': 'content'
        },
        {
            'title': '感谢聆听',
            'page_type': 'ending'
        },
    ]

    try:
        result = render_ultimate_ppt(template_path, test_slides)

        output_file = os.path.join(os.path.dirname(template_path), "v3_test_output.pptx")
        with open(output_file, 'wb') as f:
            f.write(result.getvalue())

        print(f"\n[SUCCESS] 测试完成!")
        print(f"  输出文件: {output_file}")
        print(f"  文件大小: {len(result.getvalue()) / 1024:.1f} KB")
        print(f"\n验证要点:")
        print(f"  1. 打开 PPT，应该只看到你上面的 5 页内容")
        print(f"  2. 不应有任何模板原始的示例文字或图片")
        print(f"  3. 字体、颜色、背景应与模板一致")
        print(f"  4. 每页的文字都在正确的占位符位置")

    except Exception as e:
        print(f"\n[ERROR] 测试失败: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
